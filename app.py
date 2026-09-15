"""
Gerador de boletins em PowerPoint
=================================

Recebe:
  - 1 arquivo .pptx modelo (1 slide com os marcadores {{NOME}}, {{BOOK}}, ...)
  - 1 planilha .xlsx com abas "TURMA ..." e as abas "Grade A/B/C/D"

Devolve:
  - 1 arquivo .pptx por TURMA, com um slide por aluno
  - 1 .zip com todos

Rodar:  streamlit run app.py
"""

from __future__ import annotations

import datetime as _dt
import io
import os
import posixpath
import random
import re
import zipfile
from typing import Dict, List, Tuple

import openpyxl
import streamlit as st
from pptx import Presentation

# ----------------------------------------------------------------------------
# CONFIGURAÇÕES  (mexa só aqui se a planilha mudar de colunas)
# ----------------------------------------------------------------------------

SENHA_PADRAO = "030826"

# Marcador no PPT  ->  letra da coluna na planilha
MAPA_COLUNAS: Dict[str, str] = {
    "NOME": "A",
    "LIS": "B",
    "SP": "C",
    "WR": "D",
    "PREP": "E",
    "CONS": "F",
    "ABSENCES": "H",
    "BOOK": "I",
    "MIDTERM TEST": "J",
    "CLASS": "K",
    "DATE": "L",
}

COLUNA_OVERALL = "G"          # define qual aba "Grade X" fornece a mensagem
COLUNA_NOME = "A"             # linha sem nome = linha ignorada
MARCADOR_MENSAGEM = "MESSAGE"

# Colunas que na planilha se repetem por fórmula: se vierem vazias,
# o app repete o último valor preenchido acima.
COLUNAS_REPETEM = ("I", "J", "K", "L")

# Apelidos aceitos nos marcadores (tolerante a variações de escrita)
APELIDOS: Dict[str, str] = {
    "NAME": "NOME",
    "STUDENT": "NOME",
    "ALUNO": "NOME",
    "ABSCENCES": "ABSENCES",
    "FALTAS": "ABSENCES",
    "MIDTERM": "MIDTERM TEST",
    "MENSAGEM": "MESSAGE",
    "DATA": "DATE",
    "LISTENING": "LIS",
    "SPEAKING": "SP",
    "WRITING": "WR",
}

NOTAS_EM_PONTOS = {"A": 10.0, "B": 7.5, "C": 5.0, "D": 2.5}

RE_PLACEHOLDER = re.compile(r"\{\{([^{}]*)\}\}")

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
TIPO_SLIDE = f"{NS_R}/slide"
TIPO_NOTES = f"{NS_R}/notesSlide"
CT_SLIDE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"


# ----------------------------------------------------------------------------
# 1. LEITURA DA PLANILHA
# ----------------------------------------------------------------------------

def _primeira_secao(fmt: str) -> str:
    """Devolve a 1ª seção do formato numérico (a parte usada por valores positivos)."""
    partes, atual, aspas = [], [], False
    for ch in fmt or "":
        if ch == '"':
            aspas = not aspas
            atual.append(ch)
        elif ch == ";" and not aspas:
            partes.append("".join(atual))
            atual = []
        else:
            atual.append(ch)
    partes.append("".join(atual))
    return partes[0]


def _data_para_texto(valor, fmt: str) -> str:
    """Converte data respeitando o formato do Excel (dd/mm/yyyy, dd"/"mm, ...)."""
    fmt_limpo = re.sub(r"\[[^\]]*\]", "", _primeira_secao(fmt))
    if not fmt_limpo or fmt_limpo.strip().lower() == "general":
        return valor.strftime("%d/%m/%Y")

    tokens = [
        ("yyyy", "%Y"), ("yy", "%y"),
        ("mmmm", "%B"), ("mmm", "%b"), ("mm", "%m"), ("m", "%m"),
        ("dddd", "%A"), ("ddd", "%a"), ("dd", "%d"), ("d", "%d"),
        ("hh", "%H"), ("h", "%H"), ("ss", "%S"), ("s", "%S"),
    ]
    baixo, saida, i = fmt_limpo.lower(), [], 0
    while i < len(baixo):
        ch = baixo[i]
        if ch == '"':                                   # literal entre aspas
            fim = baixo.find('"', i + 1)
            fim = len(baixo) if fim == -1 else fim
            saida.append(fmt_limpo[i + 1:fim])
            i = fim + 1
            continue
        if ch == "\\":                                  # caractere escapado
            if i + 1 < len(baixo):
                saida.append(fmt_limpo[i + 1])
            i += 2
            continue
        for token, strf in tokens:
            if baixo.startswith(token, i):
                saida.append(strf)
                i += len(token)
                break
        else:
            saida.append(fmt_limpo[i])
            i += 1
    try:
        return valor.strftime("".join(saida))
    except Exception:
        return valor.strftime("%d/%m/%Y")


def _numero_para_texto(valor: float, fmt: str) -> str:
    """Formata número como o Excel mostra: 4 com formato 0"/7" vira 4/7."""
    fmt_limpo = re.sub(r"\[[^\]]*\]", "", _primeira_secao(fmt))

    if not fmt_limpo or fmt_limpo.strip().lower() == "general":
        if float(valor).is_integer():
            return str(int(valor))
        return ("%g" % valor).replace(".", ",")

    antes, depois, padrao = [], [], []
    comecou = terminou = False
    i = 0
    while i < len(fmt_limpo):
        ch = fmt_limpo[i]
        if ch == '"':
            fim = fmt_limpo.find('"', i + 1)
            fim = len(fmt_limpo) if fim == -1 else fim
            (depois if comecou else antes).append(fmt_limpo[i + 1:fim])
            i = fim + 1
            continue
        if ch == "\\":
            if i + 1 < len(fmt_limpo):
                (depois if comecou else antes).append(fmt_limpo[i + 1])
            i += 2
            continue
        if ch in "0#?,." and not terminou:
            padrao.append(ch)
            comecou = True
            i += 1
            continue
        if comecou:
            terminou = True
        (depois if comecou else antes).append(ch)
        i += 1

    padrao_txt, prefixo, sufixo = "".join(padrao), "".join(antes), "".join(depois)
    numero = float(valor)
    if "%" in prefixo + sufixo:
        numero *= 100

    casas = len(padrao_txt.split(".")[-1].replace(",", "")) if "." in padrao_txt else 0
    milhar = "," in padrao_txt.split(".")[0]

    if milhar:
        corpo = f"{numero:,.{casas}f}"
        corpo = corpo.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
    else:
        corpo = f"{numero:.{casas}f}"
        if casas:
            corpo = corpo.replace(".", ",")
    return f"{prefixo}{corpo}{sufixo}"


def celula_para_texto(valor, number_format: str = "General") -> str:
    """Transforma o conteúdo de uma célula no texto que o Excel exibe na tela."""
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor.strip()
    if isinstance(valor, bool):
        return "SIM" if valor else "NAO"
    if isinstance(valor, (_dt.datetime, _dt.date)):
        return _data_para_texto(valor, number_format)
    if isinstance(valor, _dt.time):
        return valor.strftime("%H:%M")
    if isinstance(valor, (int, float)):
        return _numero_para_texto(float(valor), number_format)
    return str(valor).strip()


def _overall_calculado(lis: str, sp: str, wr: str) -> str:
    """Recalcula o conceito geral quando a planilha não trouxe o valor da fórmula."""
    notas = [str(x).strip().upper() for x in (lis, sp, wr)]
    if not all(n in NOTAS_EM_PONTOS for n in notas):
        return ""
    pct = sum(NOTAS_EM_PONTOS[n] for n in notas) / 30
    if pct >= 0.85:
        return "A"
    if pct >= 0.70:
        return "B"
    if pct >= 0.50:
        return "C"
    return "D"


def ler_planilha(arquivo_bytes: bytes):
    """Lê a planilha e devolve (turmas, mensagens_por_grade, avisos)."""
    avisos: List[str] = []
    wb = openpyxl.load_workbook(io.BytesIO(arquivo_bytes), data_only=True)

    # --- mensagens das abas Grade A/B/C/D --------------------------------
    mensagens: Dict[str, List[str]] = {}
    for aba in wb.sheetnames:
        achou = re.fullmatch(r"\s*grade\s*([A-D])\s*", aba, flags=re.IGNORECASE)
        if not achou:
            continue
        letra = achou.group(1).upper()
        textos: List[str] = []
        for (celula,) in wb[aba].iter_rows(min_col=1, max_col=1, values_only=True):
            texto = "" if celula is None else str(celula).strip()
            if not texto:
                continue
            if not textos and texto.lower() in {
                "mensagem", "mensagens", "message", "messages",
                letra.lower(), f"grade {letra.lower()}",
            }:
                continue                                  # cabeçalho
            textos.append(texto)
        mensagens[letra] = textos
        if not textos:
            avisos.append(f'A aba "{aba}" está sem mensagens.')

    # --- alunos das abas TURMA ------------------------------------------
    turmas: Dict[str, List[dict]] = {}
    for aba in wb.sheetnames:
        if not re.match(r"\s*turma\b", aba, flags=re.IGNORECASE):
            continue
        ws = wb[aba]
        linhas: List[dict] = []
        ultimo: Dict[str, Tuple[object, str]] = {}

        for idx in range(2, (ws.max_row or 1) + 1):       # linha 1 = cabeçalho
            nome = ws[f"{COLUNA_NOME}{idx}"].value
            if nome is None or not str(nome).strip():
                continue

            valores: Dict[str, str] = {}
            for marcador, col in MAPA_COLUNAS.items():
                cel = ws[f"{col}{idx}"]
                valor, fmt = cel.value, cel.number_format
                vazio = valor is None or str(valor).strip() == ""
                if col in COLUNAS_REPETEM:
                    if vazio:
                        valor, fmt = ultimo.get(col, (None, fmt))     # repete de cima
                    else:
                        ultimo[col] = (valor, fmt)
                elif col == "H" and vazio:
                    valor, fmt = 0, "0"                              # sem faltas = 0
                valores[marcador] = celula_para_texto(valor, fmt)

            overall = celula_para_texto(ws[f"{COLUNA_OVERALL}{idx}"].value).upper()
            if overall not in NOTAS_EM_PONTOS:
                overall = _overall_calculado(
                    valores.get("LIS", ""), valores.get("SP", ""), valores.get("WR", "")
                )
                if overall:
                    avisos.append(
                        f'{aba}, linha {idx}: coluna "Overall" veio vazia; conceito '
                        f"{overall} calculado a partir de Listening/Speaking/Writing."
                    )
            if overall not in NOTAS_EM_PONTOS:
                avisos.append(
                    f'{aba}, linha {idx} ({valores.get("NOME", "?")}): sem conceito em '
                    '"Overall" — o slide ficará sem mensagem.'
                )

            linhas.append({"linha": idx, "valores": valores, "overall": overall})

        if linhas:
            turmas[aba] = linhas
        else:
            avisos.append(f'A aba "{aba}" não tem alunos preenchidos.')

    return turmas, mensagens, avisos


# ----------------------------------------------------------------------------
# 2. SORTEIO DAS MENSAGENS
# ----------------------------------------------------------------------------

class SorteadorDeMensagens:
    """Sorteia mensagens sem repetir enquanto ainda houver opções no grupo."""

    def __init__(self, mensagens: Dict[str, List[str]], seed=None):
        self._mensagens = mensagens
        self._rng = random.Random(seed)
        self._baralho: Dict[str, List[str]] = {}

    def sortear(self, grade: str) -> str:
        pool = self._mensagens.get(grade, [])
        if not pool:
            return ""
        if not self._baralho.get(grade):
            baralho = list(pool)
            self._rng.shuffle(baralho)
            self._baralho[grade] = baralho
        return self._baralho[grade].pop()


# ----------------------------------------------------------------------------
# 3. MONTAGEM DO POWERPOINT
# ----------------------------------------------------------------------------

def _chave(bruto: str) -> str:
    valor = re.sub(r"[\s_]+", " ", bruto).strip().upper()
    return APELIDOS.get(valor, valor)


def _escapar_xml(texto: str) -> str:
    texto = texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"[\r\n\t]+", " ", texto)


def _todas_as_formas(formas):
    for shape in formas:
        yield shape
        if hasattr(shape, "shapes"):                      # grupo de formas
            try:
                yield from _todas_as_formas(shape.shapes)
            except Exception:
                pass


def _juntar_runs(paragrafo) -> None:
    """Se um marcador estiver quebrado em vários trechos, junta tudo no primeiro."""
    runs = paragrafo.runs
    if len(runs) < 2:
        return
    inteiro = "".join(r.text for r in runs)
    marcadores = ["{{" + m + "}}" for m in RE_PLACEHOLDER.findall(inteiro)]
    if not marcadores:
        return
    if all(any(marc in r.text for r in runs) for marc in marcadores):
        return                                            # já estão inteiros
    runs[0].text = inteiro
    for r in runs[1:]:
        r.text = ""


def normalizar_modelo(pptx_bytes: bytes) -> bytes:
    """Prepara o modelo para a substituição de texto funcionar sempre."""
    prs = Presentation(io.BytesIO(pptx_bytes))
    for slide in prs.slides:
        for shape in _todas_as_formas(slide.shapes):
            if shape.has_text_frame:
                for par in shape.text_frame.paragraphs:
                    _juntar_runs(par)
            if getattr(shape, "has_table", False):
                for linha in shape.table.rows:
                    for celula in linha.cells:
                        for par in celula.text_frame.paragraphs:
                            _juntar_runs(par)
    saida = io.BytesIO()
    prs.save(saida)
    return saida.getvalue()


def _achar_slide_modelo(zin: zipfile.ZipFile) -> Tuple[str, str]:
    """Descobre qual parte do .pptx é o primeiro slide."""
    pres = zin.read("ppt/presentation.xml").decode("utf-8")
    rels = zin.read("ppt/_rels/presentation.xml.rels").decode("utf-8")

    alvos = {}
    for rel in re.findall(r"<Relationship\b[^>]*/>", rels):
        rid = re.search(r'Id="([^"]+)"', rel)
        tipo = re.search(r'Type="([^"]+)"', rel)
        alvo = re.search(r'Target="([^"]+)"', rel)
        if rid and tipo and alvo and tipo.group(1) == TIPO_SLIDE:
            alvos[rid.group(1)] = alvo.group(1)

    ordem = re.findall(r'<p:sldId\b[^>]*r:id="([^"]+)"', pres)
    if not ordem or ordem[0] not in alvos:
        raise ValueError("Não encontrei nenhum slide no PowerPoint enviado.")

    alvo = alvos[ordem[0]].lstrip("/")
    parte = alvo if alvo.startswith("ppt/") else posixpath.normpath(posixpath.join("ppt", alvo))
    rels_parte = posixpath.join(
        posixpath.dirname(parte), "_rels", posixpath.basename(parte) + ".rels"
    )
    return parte, rels_parte


def marcadores_do_modelo(pptx_bytes: bytes) -> List[str]:
    """Lista os marcadores {{...}} presentes no slide modelo."""
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as zin:
        parte, _ = _achar_slide_modelo(zin)
        xml = zin.read(parte).decode("utf-8")
    vistos, saida = set(), []
    for bruto in RE_PLACEHOLDER.findall(xml):
        chave = _chave(bruto)
        if chave not in vistos:
            vistos.add(chave)
            saida.append(chave)
    return saida


def _preencher(xml: str, ctx: Dict[str, str]) -> str:
    def troca(m: re.Match) -> str:
        chave = _chave(m.group(1))
        return _escapar_xml(ctx[chave]) if chave in ctx else m.group(0)
    return RE_PLACEHOLDER.sub(troca, xml)


def gerar_pptx(modelo_bytes: bytes, contextos: List[Dict[str, str]]) -> bytes:
    """Cria um .pptx com um slide por item de `contextos`, copiando o slide modelo."""
    if not contextos:
        raise ValueError("Nenhum aluno para gerar slides.")

    with zipfile.ZipFile(io.BytesIO(modelo_bytes)) as zin:
        nomes = set(zin.namelist())
        parte_slide, parte_rels = _achar_slide_modelo(zin)
        xml_slide = zin.read(parte_slide).decode("utf-8")
        xml_rels = (
            zin.read(parte_rels).decode("utf-8") if parte_rels in nomes
            else '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/'
                 'package/2006/relationships"/>'
        )
        pres = zin.read("ppt/presentation.xml").decode("utf-8")
        pres_rels = zin.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
        outras_partes = {
            nome: zin.read(nome)
            for nome in zin.namelist()
            if not nome.startswith("ppt/slides/")
            and not nome.startswith("ppt/notesSlides/")
            and nome not in {
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
                "[Content_Types].xml",
            }
        }

    total = len(contextos)

    # rels de cada slide novo: mantém layout e imagens, tira as notas do modelo
    rels_slide = re.sub(
        r'<Relationship\b[^>]*Type="%s"[^>]*/>' % re.escape(TIPO_NOTES), "", xml_rels
    )

    # presentation.xml.rels: troca as relações de slide pelas novas
    rels_base = re.sub(
        r'<Relationship\b[^>]*Type="%s"[^>]*/>' % re.escape(TIPO_SLIDE), "", pres_rels
    )
    novas_rels = "".join(
        f'<Relationship Id="rIdSlide{i}" Type="{TIPO_SLIDE}" Target="slides/slide{i}.xml"/>'
        for i in range(1, total + 1)
    )
    pres_rels_final = rels_base.replace("</Relationships>", novas_rels + "</Relationships>")

    # presentation.xml: reescreve a lista de slides
    lista = "".join(
        f'<p:sldId id="{255 + i}" r:id="rIdSlide{i}"/>' for i in range(1, total + 1)
    )
    if re.search(r"<p:sldIdLst\b.*?</p:sldIdLst>", pres, flags=re.S):
        pres_final = re.sub(
            r"<p:sldIdLst\b.*?</p:sldIdLst>",
            f"<p:sldIdLst>{lista}</p:sldIdLst>",
            pres, count=1, flags=re.S,
        )
    elif "<p:sldIdLst/>" in pres:
        pres_final = pres.replace("<p:sldIdLst/>", f"<p:sldIdLst>{lista}</p:sldIdLst>", 1)
    else:
        pres_final = pres.replace(
            "</p:sldMasterIdLst>",
            f"</p:sldMasterIdLst><p:sldIdLst>{lista}</p:sldIdLst>", 1,
        )

    # [Content_Types].xml: um Override por slide novo
    ct = re.sub(r'<Override\b[^>]*PartName="/ppt/slides/[^"]*"[^>]*/>', "", content_types)
    ct = re.sub(r'<Override\b[^>]*PartName="/ppt/notesSlides/[^"]*"[^>]*/>', "", ct)
    novos_ct = "".join(
        f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="{CT_SLIDE}"/>'
        for i in range(1, total + 1)
    )
    ct_final = ct.replace("</Types>", novos_ct + "</Types>")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
        for nome, dados in outras_partes.items():
            zout.writestr(nome, dados)
        zout.writestr("[Content_Types].xml", ct_final)
        zout.writestr("ppt/presentation.xml", pres_final)
        zout.writestr("ppt/_rels/presentation.xml.rels", pres_rels_final)
        for i, ctx in enumerate(contextos, start=1):
            zout.writestr(f"ppt/slides/slide{i}.xml", _preencher(xml_slide, ctx))
            zout.writestr(f"ppt/slides/_rels/slide{i}.xml.rels", rels_slide)
    return buffer.getvalue()


# ----------------------------------------------------------------------------
# 4. ORQUESTRAÇÃO
# ----------------------------------------------------------------------------

def _nome_arquivo(turma: str, prefixo: str = "") -> str:
    base = re.sub(r"[^\w\s-]", "", f"{prefixo}{turma}").strip().replace(" ", "_")
    return f"{base or 'TURMA'}.pptx"


def gerar_tudo(modelo_bytes: bytes, planilha_bytes: bytes, seed=None, prefixo=""):
    turmas, mensagens, avisos = ler_planilha(planilha_bytes)
    modelo = normalizar_modelo(modelo_bytes)
    sorteador = SorteadorDeMensagens(mensagens, seed=seed)

    resultados = []
    for turma, linhas in turmas.items():
        contextos = []
        for item in linhas:
            ctx = dict(item["valores"])
            grade = item["overall"]
            ctx[MARCADOR_MENSAGEM] = sorteador.sortear(grade) if grade else ""
            if grade and not ctx[MARCADOR_MENSAGEM]:
                avisos.append(
                    f'Sem mensagens na aba "Grade {grade}" para {ctx.get("NOME", "?")}.'
                )
            contextos.append(ctx)
        resultados.append({
            "turma": turma,
            "arquivo": _nome_arquivo(turma, prefixo),
            "bytes": gerar_pptx(modelo, contextos),
            "slides": len(contextos),
            "alunos": [c.get("NOME", "") for c in contextos],
        })
    return resultados, avisos


def zipar(resultados) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        for r in resultados:
            z.writestr(r["arquivo"], r["bytes"])
    return buffer.getvalue()


# ----------------------------------------------------------------------------
# 5. INTERFACE STREAMLIT
# ----------------------------------------------------------------------------

def _senha_ok() -> bool:
    esperada = os.environ.get("APP_PASSWORD") or SENHA_PADRAO
    try:
        esperada = st.secrets.get("APP_PASSWORD", esperada)
    except Exception:
        pass

    if st.session_state.get("autenticado"):
        return True

    st.title("🔒 Gerador de boletins")
    senha = st.text_input("Senha", type="password")
    if st.button("Entrar", type="primary"):
        if senha == str(esperada):
            st.session_state["autenticado"] = True
            st.rerun()
        else:
            st.error("Senha incorreta.")
    return False


def main() -> None:
    st.set_page_config(page_title="Gerador de boletins", page_icon="📊", layout="centered")

    if not _senha_ok():
        st.stop()

    st.title("📊 Gerador de boletins em PowerPoint")
    st.caption("Envie o modelo .pptx e a planilha .xlsx. Sai um PowerPoint por turma.")

    col1, col2 = st.columns(2)
    with col1:
        modelo_up = st.file_uploader("Modelo PowerPoint (.pptx)", type=["pptx"])
    with col2:
        planilha_up = st.file_uploader("Planilha (.xlsx)", type=["xlsx", "xlsm"])

    with st.expander("Opções"):
        prefixo = st.text_input("Prefixo do nome dos arquivos", value="Boletim_")
        fixar = st.checkbox(
            "Fixar o sorteio (a mesma planilha gera sempre as mesmas mensagens)",
            value=False,
        )
        seed = int(st.number_input("Semente do sorteio", value=42, step=1)) if fixar else None

    if not (modelo_up and planilha_up):
        st.info("Envie os dois arquivos para continuar.")
        return

    modelo_bytes = modelo_up.getvalue()
    planilha_bytes = planilha_up.getvalue()

    try:
        turmas, mensagens, avisos = ler_planilha(planilha_bytes)
    except Exception as erro:
        st.error(f"Não consegui ler a planilha: {erro}")
        return

    if not turmas:
        st.error('Nenhuma aba começando com "TURMA" foi encontrada na planilha.')
        return

    st.subheader("Prévia")
    st.table([{"Turma": t, "Alunos (slides)": len(l)} for t, l in turmas.items()])
    if mensagens:
        st.write(
            "Mensagens disponíveis — "
            + " | ".join(f"Grade {g}: {len(m)}" for g, m in sorted(mensagens.items()))
        )

    try:
        encontrados = marcadores_do_modelo(modelo_bytes)
    except Exception as erro:
        st.error(f"Não consegui ler o PowerPoint: {erro}")
        return

    conhecidos = set(MAPA_COLUNAS) | {MARCADOR_MENSAGEM}
    desconhecidos = [m for m in encontrados if m not in conhecidos]
    faltando = [m for m in conhecidos if m not in encontrados]
    if desconhecidos:
        st.warning("Marcadores do modelo que o app não conhece: " + ", ".join(desconhecidos))
    if faltando:
        st.info("Marcadores previstos que não estão no modelo: " + ", ".join(sorted(faltando)))
    for aviso in avisos:
        st.warning(aviso)

    if st.button("Gerar PowerPoints", type="primary"):
        with st.spinner("Gerando..."):
            try:
                resultados, _ = gerar_tudo(
                    modelo_bytes, planilha_bytes, seed=seed, prefixo=prefixo
                )
            except Exception as erro:
                st.error(f"Erro ao gerar: {erro}")
                return
        st.session_state["resultados"] = resultados

    resultados = st.session_state.get("resultados")
    if resultados:
        st.success(
            f"{len(resultados)} arquivo(s) gerado(s), "
            f"{sum(r['slides'] for r in resultados)} slides no total."
        )
        st.download_button(
            "⬇️ Baixar tudo (.zip)",
            data=zipar(resultados),
            file_name="boletins.zip",
            mime="application/zip",
            type="primary",
        )
        for r in resultados:
            st.download_button(
                f"⬇️ {r['arquivo']} ({r['slides']} slides)",
                data=r["bytes"],
                file_name=r["arquivo"],
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation"
                ),
                key=f"dl_{r['turma']}",
            )


if __name__ == "__main__":
    main()
